import io
import base64
import pypdf
import docx

class DocumentParser:
    """
    Independent document text extractor supporting PDF, DOCX, TXT, and Base64 payloads.
    Does not require any external OCR APIs or LLMs.
    """

    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, content_type: str = None, file_name: str = "") -> str:
        if not file_bytes:
            return ""

        extracted_text = ""
        lower_name = (file_name or "").lower()
        lower_type = (content_type or "").lower()

        # 1. Try PDF extraction
        if "pdf" in lower_type or lower_name.endswith(".pdf"):
            try:
                pdf_reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                text_parts = []
                for page in pdf_reader.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
                extracted_text = "\n".join(text_parts)
            except Exception as e:
                print(f"[DocumentParser] PDF extraction error: {e}")

        # 2. Try DOCX extraction
        elif "word" in lower_type or "docx" in lower_type or lower_name.endswith(".docx"):
            try:
                doc = docx.Document(io.BytesIO(file_bytes))
                text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            text_parts.append(" | ".join(row_text))
                extracted_text = "\n".join(text_parts)
            except Exception as e:
                print(f"[DocumentParser] DOCX extraction error: {e}")

        # 3. Fallback to direct text/UTF-8 decoding
        if not extracted_text.strip():
            for encoding in ("utf-8", "latin-1", "ascii", "utf-16"):
                try:
                    extracted_text = file_bytes.decode(encoding)
                    if any(c.isalnum() for c in extracted_text):
                        break
                except UnicodeDecodeError:
                    continue

        return extracted_text.strip()

    @classmethod
    def extract_text_from_base64(cls, base64_str: str, content_type: str = None, file_name: str = "") -> str:
        if not base64_str:
            return ""
        try:
            # Handle data URL prefixes if present e.g. "data:application/pdf;base64,..."
            if "," in base64_str and "base64" in base64_str:
                base64_str = base64_str.split(",", 1)[1]
            file_bytes = base64.b64decode(base64_str)
            return cls.extract_text_from_bytes(file_bytes, content_type, file_name)
        except Exception as e:
            print(f"[DocumentParser] Base64 decoding failed: {e}")
            return ""
